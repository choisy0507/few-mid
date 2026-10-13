"""Generate editable PPTX, a self-contained slide viewer, diagrams, Word and PDF-ready text.
Run from the repository root: python3 submission/source/build_submission.py
PDF and screenshot export: node submission/source/export.mjs
"""
from pathlib import Path
from html import escape
import json, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement as PptElement
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'assets'
PAPER, INK, MUTED, LINE, LIME = '#F6F5EF', '#263D32', '#657369', '#D7DACF', '#DEF19F'
SANS, SERIF = 'Noto Sans CJK KR', 'Georgia'

def svg_text(x, y, text, size=23, color=INK, anchor='start', weight=400):
    lines = text.split('\n')
    return f'<text x="{x}" y="{y}" font-size="{size}" fill="{color}" text-anchor="{anchor}" font-weight="{weight}">' + ''.join(f'<tspan x="{x}" dy="{0 if i == 0 else size * 1.5}">{escape(line)}</tspan>' for i, line in enumerate(lines)) + '</text>'

def svg_box(x,y,w,h,fill='#FBFBF8',stroke=LINE,rx=8):
    return f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{rx}" fill="{fill}" stroke="{stroke}"/>'

def arrow(x1,y1,x2,y2):
    return f'<path d="M{x1} {y1} L{x2} {y2}" stroke="#7D8F7D" stroke-width="2" fill="none" marker-end="url(#a)"/>'

def save_svg(name, body, w=1400,h=700):
    svg=f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><defs><marker id="a" markerWidth="8" markerHeight="8" refX="6" refY="3" orient="auto"><path d="M0 0 L6 3 L0 6" fill="none" stroke="#7D8F7D"/></marker></defs><style>text{{font-family:Arial,"Noto Sans CJK KR",sans-serif}}</style><rect width="100%" height="100%" fill="{PAPER}"/>{body}</svg>'
    (ASSETS / f'{name}.svg').write_text(svg)

def diagrams():
    b=svg_box(505,25,390,90,LIME)+svg_text(700,63,'HOME / STUDY',24,anchor='middle',weight=600)+svg_text(700,96,'질문 + 대표 장면',20,anchor='middle')
    b+=arrow(700,115,700,166)+svg_box(190,170,1020,190)
    b+=svg_text(220,205,'PERSPECTIVE / 같은 장면, 세 관점',22,weight=600)
    for x,title,desc in [(230,'Notes','구체적인 대상'),(585,'Feel','질감과 인상'),(940,'Family','큰 방향')]:
        b+=svg_box(x,228,230,85,'#ECEFE3')+svg_text(x+115,264,title,26,anchor='middle')+svg_text(x+115,294,desc,19,anchor='middle')
    b+=arrow(462,273,580,273)+arrow(817,273,935,273)
    b+=svg_text(700,345,'요소 집중 · 의미 · What to Notice / 선택한 사례 유지',19,anchor='middle')
    b+=arrow(700,360,700,405)+svg_box(465,408,470,85,LIME)+svg_text(700,444,'COMPARE',25,anchor='middle',weight=600)+svg_text(700,475,'같은 장미 + 다른 나무 노트 사례',20,anchor='middle')
    b+=arrow(520,493,295,555)+arrow(880,493,1095,555)
    for x,title,desc in [(65,'CONNECTIONS','노트 ↔ 표현 ↔ 계열'),(925,'ABOUT / SOURCES','읽는 방법 · 콘텐츠 기준')]:
        b+=svg_box(x,558,410,90)+svg_text(x+205,593,title,23,anchor='middle',weight=600)+svg_text(x+205,625,desc,20,anchor='middle')
    b+=svg_text(700,682,'한 페이지의 영역과 상태 / Notes·Feel·Family는 독립 페이지가 아닙니다.',17,MUTED,anchor='middle')
    save_svg('sitemap',b)

    b=''
    steps=[(30,'01 ARRIVE','질문과 장면을 만남'),(300,'02 CHOOSE','관점 선택 / 자유 왕복'),(570,'03 FOCUS','요소와 역할 관찰'),(840,'04 COMPARE','같은 부분 · 바뀐 부분'),(1110,'05 CONNECT','관련 개념으로 확장')]
    for x,t,d in steps:
        b+=svg_box(x,90,245,112,LIME if x==300 else '#FBFBF8')+svg_text(x+18,133,t,21,weight=600)+svg_text(x+18,175,d,18)
        if x<1110:b+=arrow(x+247,146,x+265,146)
    for x,title,desc in [(230,'Notes','대상을 알고 싶다'),(560,'Feel','질감을 이해하고 싶다'),(890,'Family','큰 방향을 보고 싶다')]:
        b+=svg_box(x,295,270,115,'#ECEFE3')+svg_text(x+135,340,title,27,anchor='middle')+svg_text(x+135,380,desc,20,anchor='middle')
    b+=arrow(420,202,365,288)+arrow(420,202,695,288)+arrow(420,202,1025,288)
    b+=svg_box(230,505,940,115,'#E9EDDF')+svg_text(700,550,'탐색 상태는 유지됩니다.',28,anchor='middle')+svg_text(700,592,'관점을 바꾸거나 비교해도 현재 사례와 관심 대상을 잃지 않습니다.',21,anchor='middle')
    b+=svg_text(700,678,'대표 흐름은 안내입니다. 사용자는 어떤 관점부터 시작해도 됩니다.',18,MUTED,anchor='middle')
    save_svg('user-flow',b)

    b=svg_box(50,20,1300,640,'#FCFCFA','#98A297')
    b+=svg_text(85,65,'THE ANATOMY OF SCENT',21,weight=600)+svg_text(1300,65,'Study / Connections / About',18,anchor='end')
    b+=f'<path d="M75 88H1325" stroke="{LINE}"/>'
    b+=svg_text(85,145,'같은 장미, 다른 나무. 어떤 차이에 주목할까요?',31)
    b+=svg_text(85,186,'① 질문과 사례: 전문 용어를 몰라도 탐색을 시작할 수 있는 맥락',18,MUTED)
    b+=svg_box(80,220,1240,50,'#F1F2EA')+svg_text(103,253,'SCENT STUDY 01',17)
    for x,t in [(840,'Notes'),(990,'Feel'),(1140,'Family')]:b+=svg_box(x,228,140,33,LIME if t=='Notes' else '#FFFFFF')+svg_text(x+70,252,t,17,anchor='middle')
    b+=svg_box(80,270,830,280,'#EFEFE8')+svg_box(910,270,410,280)
    b+=f'<ellipse cx="400" cy="408" rx="150" ry="90" fill="#E1D8D2" stroke="#B8ACA3"/><path d="M530 473 Q 665 315 825 437" fill="none" stroke="#A39C88" stroke-width="14"/>'
    b+=svg_box(270,330,145,48)+svg_text(342,362,'Rose',21,anchor='middle')+svg_box(675,439,155,48)+svg_text(752,471,'Cedarwood',21,anchor='middle')
    b+=svg_text(102,527,'② 선택한 요소 강조 / 다른 부분 약화',16,MUTED)
    b+=svg_text(940,310,'◇ NOTE / 정보 유형',16,MUTED)+svg_text(940,353,'이름 · 의미',27)+svg_text(940,389,'구체적인 장면과 표현의 역할',18)+svg_box(935,423,360,81,'#E9EDDF')+svg_text(950,452,'③ What to Notice',18,weight=600)+svg_text(950,482,'시향에서 구별할 질문',18)
    b+=svg_box(80,550,1240,82,'#F1F2EA')+svg_text(102,582,'④ 사례 전환: Rose는 유지, 나무 노트만 비교',19)+svg_text(1110,583,'Cedarwood ↔ Sandalwood',19,anchor='middle')+svg_text(102,612,'유지된 것 / 바뀐 것 / 구별할 단서',17,MUTED)
    b+=svg_text(700,694,'DESKTOP WIREFRAME / 위치와 위계, 행동의 목적을 먼저 설계',17,MUTED,anchor='middle')
    save_svg('wireframe-desktop',b)

    b=''
    for x,mode,label,caption in [(25,'NOTES','Rose / Cedarwood','대상별 지점과 이름'),(485,'FEEL','Dry / 꽃잎의 인상','여러 요소에 걸친 질감'),(945,'FAMILY','Floral / Woody','더 넓은 방향을 묶는 영역')]:
        b+=svg_box(x,60,430,460,'#FCFCFA','#9FAA9C')+svg_text(x+25,105,mode,26,weight=600)+svg_box(x+20,133,390,274,'#EDEFE5')
        b+=f'<ellipse cx="{x+175}" cy="247" rx="90" ry="62" fill="#DFD8D0"/><path d="M{x+170} 344Q{x+290} 185 {x+367} 300" fill="none" stroke="#A59A86" stroke-width="8"/>'
        if mode=='FAMILY':b+=f'<ellipse cx="{x+178}" cy="249" rx="130" ry="99" fill="none" stroke="#789063" stroke-dasharray="6 5"/><ellipse cx="{x+283}" cy="318" rx="115" ry="68" fill="none" stroke="#789063" stroke-dasharray="6 5"/>'
        if mode=='FEEL':b+=f'<path d="M{x+75} 190L{x+200} 390M{x+95} 190L{x+220} 390M{x+115} 190L{x+240} 390" stroke="#9D9077" stroke-width="1"/>'
        b+=svg_text(x+215,450,label,24,anchor='middle')+svg_text(x+215,491,caption,19,MUTED,anchor='middle')
    b+=svg_box(220,565,960,80,LIME)+svg_text(700,613,'장면은 유지 / 관점에 따라 읽는 정보와 강조 범위가 바뀝니다.',24,anchor='middle')
    save_svg('wireframe-perspectives',b)

    b=''
    for x,name,desc,curve in [(25,'Rose + Cedarwood','마른 결과 선명한 목재의 묘사',False),(725,'Rose + Sandalwood','둥글고 매끈한 목질감의 묘사',True)]:
        b+=svg_box(x,35,650,420,'#FCFCFA','#A0AA9D')+svg_text(x+25,80,name,29)+svg_box(x+20,110,610,215,'#EDEFE5')
        b+=f'<ellipse cx="{x+220}" cy="214" rx="115" ry="80" fill="#DED2C8" stroke="#B7A295"/>'
        for i in range(6):
            b+=f'<path d="M{x+295+i*8} 294 '+(f'C{x+270+i*9} 192 {x+570+i*7} 140 {x+540+i*6} 270' if curve else f'Q{x+420+i*8} 165 {x+560+i*5} 264')+'" fill="none" stroke="#A59576" stroke-width="2"/>'
        b+=svg_text(x+25,367,'고정: Rose    변경: '+name.split(' + ')[1],21)+svg_text(x+25,410,desc,22,MUTED)
    b+=svg_box(25,488,1350,130,LIME)+svg_text(60,533,'무엇이 같고, 무엇이 달라졌나요?',27)+svg_text(60,578,'장미의 대상을 유지하고, 나무 노트 사례의 대표적인 묘사를 비교합니다.',23)
    b+=svg_text(700,677,'교육용 설명 사례 / 실제 향의 배합 결과를 예측하지 않습니다.',18,MUTED,anchor='middle')
    save_svg('wireframe-compare',b)

    b=svg_box(35,80,810,460,'#FCFCFA','#9EAA9B')+svg_text(62,120,'WIDE / 약 1024px 이상',22,weight=600)+svg_box(60,145,760,45,'#F0F1E8')+svg_text(80,176,'Notes                 Feel                 Family',21)
    b+=svg_box(60,202,500,232,'#EDEFE5')+svg_text(310,324,'시각적 장면',29,anchor='middle')+svg_box(575,202,245,232)+svg_text(697,304,'설명 패널',24,anchor='middle')+svg_text(697,347,'관찰할 단서',20,anchor='middle')+svg_box(60,450,760,62,'#E9EDDF')+svg_text(440,490,'유지된 것 / 바뀐 것 / 구별할 단서',21,anchor='middle')
    b+=svg_box(980,20,300,630,'#FCFCFA','#53654B',25)+svg_text(1130,66,'NARROW / 390px',18,anchor='middle',weight=600)
    b+=svg_box(1000,90,260,48,'#F0F1E8')+svg_text(1130,121,'Notes / Feel / Family',18,anchor='middle')+svg_box(1000,150,260,190,'#EDEFE5')+svg_text(1130,252,'시각적 장면',23,anchor='middle')+svg_box(1000,352,260,100)+svg_text(1130,396,'설명과 관찰할 단서',19,anchor='middle')+svg_box(1000,463,260,56,LIME)+svg_text(1130,499,'사례 전환',19,anchor='middle')+svg_box(1000,531,260,93,'#E9EDDF')+svg_text(1130,567,'비교 결과를',18,anchor='middle')+svg_text(1130,596,'한 열로 읽기',18,anchor='middle')
    b+=arrow(865,300,955,300)+svg_text(440,604,'약 680px에서 쌓기 / 중간 폭은 패널과 여백부터 조정',20,MUTED,anchor='middle')+svg_text(700,691,'화면은 재배치되지만 관점·사례·집중 상태는 유지됩니다.',18,MUTED,anchor='middle')
    save_svg('wireframe-responsive',b)

slides=[]
def slide(kicker,title,subtitle='',dark=False,notes=''):
    s={'kicker':kicker,'title':title,'subtitle':subtitle,'dark':dark,'notes':notes,'elements':[]};slides.append(s)
    return s
def text(s,x,y,w,h,t,size=28,color=INK,font=SANS,bold=False,italic=False,href=None):
    s['elements'].append(dict(kind='text',x=x,y=y,w=w,h=h,text=t,size=size,color=color,font=font,bold=bold,italic=italic,href=href))
def rect(s,x,y,w,h,fill,stroke=None,radius=False):s['elements'].append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,radius=radius))
def image(s,x,y,w,h,path):s['elements'].append(dict(kind='image',x=x,y=y,w=w,h=h,path=path))
def card(s,x,y,w,h,k,title,body,fill='#ECEFE3'):
    rect(s,x,y,w,h,fill,radius=True);text(s,x+27,y+25,w-54,35,k,18,MUTED);text(s,x+27,y+80,w-54,95,title,34);text(s,x+27,y+192,w-54,h-205,body,25)

def make_slides():
    s=slide('FRONT-END WEB / WEEK 06','',notes='약 20초. 향수 설명을 읽어도 차이를 상상하기 어렵다는 경험에서 시작했습니다. 하나의 향을 여러 관점으로 읽으며 그 관계를 이해하는 사이트를 제안합니다.')
    text(s,75,110,800,300,'The Anatomy\nof Scent',108,font=SERIF);text(s,82,430,750,80,'One scent. Three perspectives.',33,font=SERIF,italic=True)
    text(s,82,525,725,155,'하나의 향 설명을 여러 관점으로 읽고,\n작은 차이를 직접 발견하는 웹 경험.',34)
    text(s,82,737,740,45,'MIDTERM DESIGN PROPOSAL · 2026',17,MUTED);image(s,925,116,590,590,'assets/moodboard.png')

    s=slide('01 / THE PROBLEM','단어를 알아도, 차이는 남습니다.','“Woody = 나무 향” 다음에 이해해야 할 것은 무엇일까요?',notes='약 25초. 단어의 정의만으로는 마른 나무 결과 둥근 목질감을 구별하기 어렵습니다. 정보의 양보다 정보를 해석하는 기준에 초점을 맞춥니다.')
    rect(s,80,280,1440,135,'#E9EDDF');text(s,115,320,1350,65,'“Powdery floral with soft musk…”',48,font=SERIF,italic=True)
    card(s,80,463,450,282,'01 / MEANING','어떤 느낌인가?','단어가 가리키는 대상과\n질감이 막연합니다.');card(s,575,463,450,282,'02 / RELATION','무엇이 다른가?','비슷한 표현의 공통점과\n구별 기준이 흐립니다.');card(s,1070,463,450,282,'03 / ATTENTION','어디에 주목할까?','다음 시향에서 관찰할\n구체적인 질문이 필요합니다.')

    s=slide('02 / CONCEPT & PREMISE','하나의 향을 읽는 세 가지 관점.','같은 장면에서 대상, 질감, 큰 방향을 번갈아 읽습니다.',notes='약 25초. Notes, Feel, Family를 바꿔 보는 경험이 핵심입니다. 장면은 유지되고 읽는 정보가 달라집니다. 실제 냄새를 재현하기보다 설명을 읽는 단서를 제공합니다.')
    text(s,83,285,680,230,'Read the scent.\nReveal its structure.',68,font=SERIF);text(s,83,565,690,100,'향 설명 → 관점 전환 → 차이 관찰 → 관계 이해',26);text(s,83,706,700,75,'대상: 용어가 익숙하지 않거나,\n단어를 알아도 조합이 막연한 방문자',22,MUTED)
    for i,(t,d) in enumerate([('Notes','무엇을 떠올리나요?'),('Feel','어떻게 느껴지나요?'),('Family','어느 방향인가요?')]):
        rect(s,880,283+i*166,625,140,LIME if i==1 else '#ECEFE3',radius=True);text(s,909,308+i*166,230,66,t,46,font=SERIF);text(s,1160,334+i*166,300,40,d,25)

    s=slide('03 / INFORMATION STRUCTURE','관련되어도, 같은 종류의 정보는 아닙니다.','한 목록에 섞지 않고, 각각 답하는 질문을 구분합니다.',notes='약 25초. Rose는 노트, Creamy는 감각 표현, Woody는 계열입니다. 연결될 수 있지만 동의어는 아닙니다. 노트가 성분표나 배합 비율을 뜻하지 않는다는 점도 설명합니다.')
    for x,k,t,body in [(80,'◇ NOTE / 대상','Rose · Sandalwood','구체적인 향의 대상.\n실제 원료의 사용이나\n배합 비율을 보장하지 않습니다.'),(575,'○ DESCRIPTOR / 질감과 인상','Dry · Creamy','느껴지는 방식의 표현.\n여러 노트와 계열에\n겹쳐 적용될 수 있습니다.'),(1070,'● FAMILY / 큰 방향','Floral · Woody','전체적인 성격을 묶는 분류.\n분류 체계에 따라\n이름과 경계가 달라집니다.')]:card(s,x,290,450,393,k,t,body)
    text(s,90,732,1390,59,'핵심: 이름을 외우는 것에서, 무엇을 설명하는지 읽는 것으로.',30)

    s=slide('04 / INTENDED EXPERIENCE','탐색 전후에 무엇이 달라질까요?',notes='약 20초. 사용자는 차분한 호기심에서 출발해 대상, 질감, 계열의 차이를 알아차리고 다음 시향의 질문을 얻습니다. 이 학습 효과는 피드백으로 검증할 계획입니다.')
    for i,(k,t,b) in enumerate([('FEEL','호기심','조용한 관찰실의\n차분한 분위기'),('NOTICE','정보의 차이','대상·질감·큰 방향의\n서로 다른 역할'),('DISCOVER','다른 가능성','같은 우디 방향 안의\n다른 질감'),('UNDERSTAND','연결의 의미','관련되어도\n동의어는 아님'),('DO','직접 관찰','선택·집중·전환·비교·\n다시 연결하기')]):
        rect(s,80+i*290,275,265,330,'#ECEFE3',radius=True);text(s,101+i*290,307,220,35,k,17,MUTED);text(s,101+i*290,377,220,50,t,32);text(s,101+i*290,470,220,95,b,22)
    rect(s,80,662,1440,140,LIME);text(s,110,687,1370,95,'“그냥 나무 향?” → “마른 결과 매끈한 질감을 따로 관찰해볼 수 있겠다.”',29)

    s=slide('05 / THE DIFFERENCE','차별화는 기능보다, 이해의 순서에 있습니다.','지도와 비교가 존재한다는 사실만으로 새로움을 주장하지 않습니다.',notes='약 20초. 분류표나 제품 목록보다 구체적인 질문을 먼저 보여줍니다. 사용자의 조작이 설명을 읽는 과정에 연결됩니다. 시장 전체에서 유일한 서비스라고 주장하지 않습니다.')
    for x,k,t,b in [(80,'ENTRY','하나의 궁금증','같은 장미, 다른 나무.\n구체적인 차이를\n탐색의 입구로 삼습니다.'),(575,'ACTION','같은 장면의 다른 읽기','대상·질감·큰 방향을\n관점 전환으로\n구조적으로 경험합니다.'),(1070,'OUTCOME','다음 시향의 단서','공통점·바뀐 요소·관찰 질문을\n돌려주어 사용자가\n차이를 말할 수 있게 합니다.')]:card(s,x,285,450,402,k,t,b)
    text(s,90,738,1400,45,'분류는 탐색의 출발점이면서 동시에, 사례를 이해한 뒤 다시 읽을 수 있는 정보입니다.',24,MUTED)

    s=slide('06 / REFERENCES & INSPIRATION','참고 자료를 설계 판단으로 연결합니다.','사용자 제공 레퍼런스 기반 분석 / 원문 접근 상태는 부록에 표시했습니다.',notes='약 25초. 향수 검색, 향 휠, 퀴즈의 정보 구조를 참고합니다. 여기에 직접 조작하는 설명과 관찰 노트의 원칙을 연결합니다. 어떤 모습보다 무엇을 배웠는지 설명합니다.')
    rows=[('Fragrantica / Perfume Finder','노트와 제품을 탐색하는 구조','제품 탐색 전에 언어와 관계를 이해'),('Scent Verdict / Scent Wheel','계열을 공간에 조직하는 접근','위치에 더해 연결의 이유를 문장으로'),('Michael Edwards / Fragrance Wheel','향 계열 분류의 참고 기준','계열·노트·표현의 역할을 구분'),('Everfumed / Scent Sign Quiz','선택으로 진입점을 만드는 경험','취향 판정 대신 다음 관찰 질문 제공')]
    for i,(a,b,c) in enumerate(rows):
        y=270+i*115;rect(s,80,y,1440,97,'#ECEFE3' if i%2==0 else '#F0F0E8');text(s,105,y+20,515,55,a,25);text(s,650,y+20,350,60,b,23);text(s,1040,y+20,450,60,c,23)
    text(s,85,770,1430,45,'경험: Explorable Explanations / 시각 언어: 에디토리얼과 식물 관찰 노트',22,MUTED)

    s=slide('07 / MOOD & VISION BOARD','Quiet. Tactile. Curious. Layered.','꽃잎, 결, 입자, 직물 — 감각을 생각해 볼 단서.',notes='약 15초. 이 프로젝트를 위해 만든 소재 콜라주입니다. 실제 향료의 증거 사진은 아니며, 서로 다른 질감을 관찰하는 분위기를 제안합니다.')
    image(s,80,270,1010,525,'assets/moodboard.png');text(s,1140,295,380,300,'조용한 관찰실\n\n만져질 듯한 소재\n\n겹침과 분리\n\n느린 발견',31);text(s,1140,691,350,98,'자체 생성 시각 자료.\n감각의 비유이며\n실제 냄새를 재현하지 않습니다.',18,MUTED)

    s=slide('08 / VISUAL LANGUAGE','읽기 위한 여백. 선택을 위한 강조.',notes='약 20초. 종이색 바탕과 깊은 초록은 읽기, 연두는 현재 선택을 위한 색입니다. 기호와 유형 이름을 함께 사용해 색만으로 정보를 구분하지 않습니다. 움직임은 선택된 대상을 알려줍니다.')
    colors=[('PAPER',PAPER),('PINE',INK),('SELECT',LIME),('ROSE','#D6A69B'),('SAND','#D8C8A5')]
    for i,(name,c) in enumerate(colors):rect(s,80+i*290,255,265,130,c,LINE);text(s,85+i*290,405,260,32,name+' / '+c,18,MUTED)
    text(s,80,485,650,150,'A word.\nA different perspective.',54,font=SERIF,italic=True);text(s,82,678,650,95,'제목: Georgia serif / 본문·UI: DM Sans + 한국어 sans\n영문 제목과 한국어 설명 사이의 위계를 명확하게.',22,MUTED)
    rect(s,860,482,660,278,'#ECEFE3',radius=True);text(s,890,512,590,45,'◇ Note    ○ Descriptor    ● Family',29);text(s,890,583,590,142,'유형 이름과 기호를 함께 표시\n선택한 요소 강조 / 다른 요소 약화\n짧은 전환 / reduced motion 지원',25)

    s=slide('09 / SITEMAP','한 페이지, 서로 연결되는 영역과 상태.',notes='약 15초. 중심 사례, 관점, 비교, 관계, 콘텐츠 기준으로 이어지는 한 페이지 구조입니다. 관점은 별도 페이지가 아니고 사용자가 자유롭게 오갈 수 있는 상태입니다.')
    image(s,155,215,1290,620,'assets/sitemap.png')

    s=slide('10 / USER FLOW','순서가 있는 안내, 자유롭게 왕복하는 탐색.',notes='약 15초. 대표 여정은 대상부터 질감과 비교로 이어집니다. 그러나 Feel이나 Family부터 시작할 수 있고 전환할 때 선택한 사례를 유지합니다.')
    image(s,120,235,1360,590,'assets/user-flow.png')

    s=slide('11 / WIREFRAME','배치의 목적을 먼저 설계합니다.','질문 → 관점 → 장면과 설명 → 하나의 차이 관찰',notes='약 20초. 질문은 왜 탐색하는지, 세 관점은 무엇을 읽는지, 장면은 어디를 보는지 알려줍니다. 설명 옆의 관찰 질문과 하단 비교가 다음 행동을 연결합니다.')
    image(s,150,245,1300,585,'assets/wireframe-desktop.png')

    s=slide('12 / KEY DESIGN STATES','장면은 유지하고, 읽는 정보가 바뀝니다.','Notes는 대상, Feel은 질감, Family는 큰 방향을 강조합니다.',notes='약 25초. 탭을 누를 때 설명 카드만 바뀌는 것이 아니라 강조의 단위가 달라집니다. 개별 대상, 여러 요소를 가로지르는 질감, 큰 방향을 묶는 영역이라는 차이를 보이게 합니다.')
    image(s,100,235,1400,560,'assets/wireframe-perspectives.png')

    s=slide('13 / DESIGN PROTOTYPE','같은 장면을 직접 바꾸어 읽어보세요.','별도 제출용 프로토타입 / 현재 실행되는 v1과 구분합니다.',notes='약 45초. 별도 HTML 프로토타입을 열어 Notes→Feel→Family를 전환하고 나무 노트 사례를 바꿉니다. 하나의 요소를 선택하면 주변이 물러납니다. 기존 첫 화면 전체를 교체한 것은 아니고 새 방향을 시험한 디자인 실험입니다.')
    image(s,80,263,1120,495,'assets/study-feel.png');text(s,1240,281,285,95,'TRY THE LENSES',17,MUTED);text(s,1240,340,280,310,'01 Notes\n대상을 읽기\n\n02 Feel\n질감을 읽기\n\n03 Family\n큰 방향 읽기',25);text(s,90,789,1420,45,'관점을 바꾸어도 현재 사례와 관심 대상은 유지됩니다.',23)

    s=slide('14 / RESPONSIVE INTENTIONS','줄이기보다, 읽는 관계를 재배치합니다.','넓은 화면은 공간적으로, 좁은 화면은 이해의 순서로.',notes='약 20초. 장면과 설명을 나란히 보다가 좁아지면 쌓습니다. 관점·사례·선택은 유지합니다. hover에 필수 정보를 숨기지 않고 키보드와 터치로 같은 정보를 읽습니다.')
    image(s,120,239,1360,559,'assets/wireframe-responsive.png');text(s,125,803,1350,37,'대략 680px에서 쌓기 · 방향키 탭 전환 · 명시적 선택 상태 · 움직임 감소 설정',20,MUTED)

    s=slide('APPENDIX A / COMPARISON DETAIL','같은 부분을 남겨야, 차이가 읽힙니다.','대표 사례 A / B — 향의 결과를 계산한 시뮬레이션이 아닙니다.',notes='질문이 나오면 사용합니다. 노트 하나를 바꾸는 비교에서 장미의 형태와 위치가 남습니다. 두 대표 묘사를 관찰하는 교육용 사례이며 배합 결과의 예측이 아닙니다.')
    image(s,90,243,1420,572,'assets/wireframe-compare.png')

    s=slide('APPENDIX B / WHAT EXISTS & WHAT COMES NEXT','현재 실험과 다음 개발 범위를 구분합니다.',notes='기존 웹사이트의 관계 지도는 구현되어 있고 11개 테스트를 통과했습니다. 제출용 관점 전환 프로토타입도 별도로 실행됩니다. 학습 효과와 향 콘텐츠는 아직 사용자·전문가 검토가 필요합니다.')
    card(s,80,276,450,437,'V1 / IMPLEMENTED','관계 지도와 비교','12개 개념 / 설명이 있는 관계\n관련 개념 강조와 표현 비교\n11개 브라우저 테스트 통과\n320–1440px 레이아웃 검증')
    card(s,575,276,450,437,'V2 / DESIGN PROTOTYPE','세 관점의 한 장면','Notes / Feel / Family 전환\n요소 집중과 사례 A/B 전환\n모바일·키보드 조작 실험\n기존 v1과 별도 파일')
    card(s,1070,276,450,437,'NEXT / REVIEW & BUILD','피드백 뒤 통합','대표 화면과 관계 지도 연결\n콘텐츠 용례와 출처 검토\n사용자의 이해 변화 확인\n시간 변화는 후속 확장')
    text(s,90,752,1420,47,'현재 오각형은 편집된 예시입니다. 다음 중심 경험은 수치보다 관찰 단서를 우선합니다.',23,MUTED)

    s=slide('APPENDIX C / COURSE CONNECTION','상호작용에는 이해를 돕는 이유가 있습니다.',notes='과제 요구사항에 연결합니다. 이야기, 경험, 시각 언어, 공간, 반응이라는 수업 주제가 같은 대표 화면에 함께 들어갑니다. 기술을 모두 보여주기보다 필요한 반응을 선택했습니다.')
    for i,(week,title,reason) in enumerate([('W01','Storytelling','추상적인 장면이 읽을 수 있는 구조로 바뀜'),('W02','Intended Experience','사용자가 질문을 갖고 선택하며 발견'),('W03','Visual Language','결·입자·색·여백이 관찰의 단서가 됨'),('W04','Spatial Relationships','대상과 역할을 공간에서 분리하고 연결'),('W05','A Web That Responds','선택과 화면 크기에 따라 강조와 배치가 변화')]):
        y=255+i*103;rect(s,80,y,1440,87,'#ECEFE3' if i%2==0 else '#F1F1E9');text(s,110,y+21,120,45,week,24,MUTED);text(s,245,y+21,445,48,title,28,font=SERIF);text(s,720,y+22,760,47,reason,25)
    text(s,85,790,1400,40,'Field notes: 관계의 이유, 좁은 화면의 재구성, 키보드·움직임 감소 실험을 기록했습니다.',20,MUTED)

    s=slide('APPENDIX D / REFERENCES & CREDIT','무엇을 참고했고, 무엇을 직접 만들었는가?',notes='외부 사이트는 네트워크 제한으로 새로 원문 확인을 못 했습니다. 제공된 대화와 URL 기반의 분석을 원문 조사와 구분했습니다. 무드보드는 자체 생성, 다이어그램은 직접 설계, 화면은 실제 프로토타입 캡처입니다.')
    refs=[('Fragrantica · Perfume Finder','https://www.fragrantica.com/perfume-finder/'),('Scent Verdict · Scent Wheel','https://scentverdict.com/guides/scentwheel'),('Michael Edwards · Fragrances of the World','https://www.fragrancesoftheworld.com/'),('Everfumed · Scent Sign Quiz','https://everfumed.com/find-your-scent-sign-quiz/'),('Bret Victor · Explorable Explanations','https://worrydream.com/ExplorableExplanations/'),('Kinfolk · Editorial Reference','https://www.kinfolk.com/')]
    for i,(label,url) in enumerate(refs):text(s,82,257+i*66,1370,52,label+'  ↗',24,href=url)
    rect(s,80,681,1440,135,'#E9EDDF');text(s,105,700,1380,94,'검토 상태: 제공 자료 기반 분석 / 외부 원문 접근은 프록시 403으로 제한.\n자체 제작: AI 생성 무드보드 · 직접 설계한 다이어그램 · 실제 프로토타입 캡처.\n시장 유일성이나 측정된 향 데이터를 주장하지 않습니다.',21)

    s=slide('APPENDIX E / FEEDBACK','탐색 뒤, 어떤 차이를 말할 수 있나요?',notes='피드백에서 새로 배운 말의 수보다 차이를 관찰하는 기준을 확인하겠습니다. 관점 전환이 실제로 도움 되는지, 이미지를 측정값으로 오해하지 않는지 물어봅니다.')
    for i,t in enumerate(['Rose · Floral · Creamy는 서로 어떤 정보인가요?','관점을 바꾸며 무엇을 새롭게 발견했나요?','두 사례에서 무엇이 유지되고 달라졌나요?','다음 시향에서 어떤 차이에 주목하겠어요?']):
        text(s,90,270+i*107,90,60,str(i+1).zfill(2),33,MUTED,font=SERIF);text(s,215,270+i*107,1280,69,t,32)
    rect(s,80,731,1440,75,LIME);text(s,115,751,1350,45,'목표: 향을 읽고, 스스로 관찰할 질문을 갖게 하는 것.',29)

    s=slide('APPENDIX F / REQUIREMENTS CHECK','제출 항목을 빠짐없이 연결했습니다.',notes='요구사항별로 발표 자료와 상세 문서, 시각 자료를 대응시킨 체크리스트입니다. 발표는 1–15장 중심으로 약 5–6분, 16–21장은 부록으로 사용합니다.')
    reqs=[('Concept / Premise','3장 / PROJECT-PROPOSAL.md'),('Intended Experience','5장 / Feel·Notice·Discover·Understand·Do'),('References / Inspiration','7장·19장 / REFERENCE-NOTES.md'),('Mood / Vision Board','8장 / assets/moodboard.png'),('Visual Language','9장 / 색·서체·기호·강조·움직임'),('Sitemap','10장 / assets/sitemap.svg'),('User Flow','11장 / assets/user-flow.svg'),('Wireframes','12–13장·16장 / 주요 상태 배치도'),('Responsive Intentions','15장 / 데스크톱·모바일 비교')]
    for i,(a,b) in enumerate(reqs):
        y=242+i*62;rect(s,80,y,1440,53,'#ECEFE3' if i%2==0 else '#F0F0E8');text(s,100,y+9,580,38,a,22);text(s,710,y+10,770,36,b,21,MUTED)

def export_slides():
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    html=[]
    for index,s in enumerate(slides,1):
        page=prs.slides.add_slide(prs.slide_layouts[6]);bg=page.background.fill;bg.solid();bg.fore_color.rgb=RGBColor.from_string(PAPER[1:])
        els=[dict(kind='text',x=80,y=45,w=1350,h=38,text=s['kicker'],size=16,color=MUTED,font=SANS,bold=False,italic=False,href=None)]
        if s['title']:els += [dict(kind='text',x=80,y=111,w=1440,h=85,text=s['title'],size=49,color=INK,font=SANS,bold=False,italic=False,href=None)]
        if s['subtitle']:els += [dict(kind='text',x=83,y=205,w=1430,h=43,text=s['subtitle'],size=23,color=MUTED,font=SANS,bold=False,italic=False,href=None)]
        els+=s['elements'];els += [dict(kind='rect',x=80,y=855,w=1440,h=1,fill=LINE,stroke=None,radius=False),dict(kind='text',x=82,y=869,w=1280,h=23,text='THE ANATOMY OF SCENT / MIDTERM DESIGN PRESENTATION',size=12,color=MUTED,font=SANS,bold=False,italic=False,href=None),dict(kind='text',x=1458,y=866,w=60,h=28,text=f'{index:02d}',size=16,color=MUTED,font=SANS,bold=False,italic=False,href=None)]
        parts=[]
        for e in els:
            x,y,w,h=[e[k] for k in ('x','y','w','h')];style=f'left:{x}px;top:{y}px;width:{w}px;height:{h}px;'
            if e['kind']=='rect':
                shape=page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['radius'] else MSO_SHAPE.RECTANGLE,Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(e['fill'][1:]);shape.line.fill.background()
                shape._element.spPr.append(PptElement('a:effectLst'))
                if e['stroke']:shape.line.color.rgb=RGBColor.from_string(e['stroke'][1:]);shape.line.width=Pt(.6)
                if e['radius']:
                    try:shape.adjustments[0]=.03
                    except Exception:pass
                parts.append(f'<div class="el" style="{style}background:{e["fill"]};border:{"1px solid "+e["stroke"] if e["stroke"] else "0"};border-radius:{"8px" if e["radius"] else "0"}"></div>')
            elif e['kind']=='image':
                # The entire image is placed without editing/cropping the generated mood board.
                from PIL import Image
                path=ROOT/e['path']; iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
                page.shapes.add_picture(str(path),Inches((x+(w-pw)/2)/120),Inches((y+(h-ph)/2)/120),width=Inches(pw/120),height=Inches(ph/120))
                parts.append(f'<img class="el" src="{escape(e["path"])}" alt="{escape(Path(e["path"]).stem)}" style="{style}object-fit:contain"/>')
            else:
                tb=page.shapes.add_textbox(Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120));tf=tb.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
                for i,line in enumerate(e['text'].split('\n')):
                    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.line_spacing=1.3;p.space_after=Pt(0);r=p.add_run();r.text=line;r.font.name=e['font'];r.font.size=Pt(e['size']*.6);r.font.bold=e['bold'];r.font.italic=e['italic'];r.font.color.rgb=RGBColor.from_string(e['color'][1:]);r._r.get_or_add_rPr().set('lang','ko-KR')
                    if e.get('href'):r.hyperlink.address=e['href']
                txt=escape(e['text']).replace('\n','<br>');txt=f'<a href="{escape(e["href"])}" target="_blank" rel="noopener">{txt}</a>' if e.get('href') else txt
                parts.append(f'<div class="el txt" style="{style}font-size:{e["size"]}px;color:{e["color"]};font-family:{escape(e["font"])},serif;font-weight:{600 if e["bold"] else 400};font-style:{"italic" if e["italic"] else "normal"}">{txt}</div>')
        page.notes_slide.notes_text_frame.text=s['notes']
        html.append(f'<section class="slide" aria-label="슬라이드 {index}">{"".join(parts)}</section>')
    prs.save(ROOT/'The-Anatomy-of-Scent.pptx')
    (ROOT/'slide-content.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2))
    css='''*{box-sizing:border-box}body{margin:0;background:#202d25;font-family:"Noto Sans CJK KR",sans-serif}.slide{position:absolute;width:1600px;height:900px;background:#F6F5EF;transform-origin:top left;display:none;overflow:hidden}.slide.active{display:block}.el{position:absolute}.txt{line-height:1.3;white-space:normal;word-break:keep-all;overflow:visible}a{color:inherit;text-decoration:none}#stage{margin:auto;position:relative}.controls{position:fixed;bottom:10px;left:50%;transform:translateX(-50%);padding:7px 15px;background:#263d32e8;color:#f6f5ef;border-radius:25px;display:flex;gap:14px;align-items:center;font-size:12px}.controls button{background:none;color:inherit;border:0;cursor:pointer;padding:6px;font:inherit}.controls button:focus-visible{outline:2px solid #def19f}.controls span{min-width:75px;text-align:center}@page{size:1600px 900px;margin:0}@media print{body{background:white}#stage{width:1600px!important;height:auto!important}.slide{display:block!important;position:relative;transform:none!important;break-after:page;page-break-after:always}.slide:last-child{break-after:auto;page-break-after:auto}.controls{display:none}}'''
    js='''const ss=[...document.querySelectorAll('.slide')];let i=0;function resize(){const scale=Math.min(innerWidth/1600,(innerHeight-25)/900);ss.forEach(s=>s.style.transform=`scale(${scale})`);document.querySelector('#stage').style.width=(1600*scale)+'px';document.querySelector('#stage').style.height=(900*scale)+'px'}function show(n){i=Math.max(0,Math.min(ss.length-1,n));ss.forEach((s,j)=>s.classList.toggle('active',j===i));document.querySelector('#count').textContent=(i+1)+' / '+ss.length;location.hash='slide-'+(i+1)}document.querySelector('#previous').onclick=()=>show(i-1);document.querySelector('#next').onclick=()=>show(i+1);document.querySelector('#print').onclick=()=>print();addEventListener('resize',resize);addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();show(i+1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(i-1)}if(e.key==='Home')show(0);if(e.key==='End')show(ss.length-1)});resize();const n=parseInt(location.hash.replace('#slide-',''));show(Number.isFinite(n)?n-1:0);window.deckReady=true;'''
    (ROOT/'presentation.html').write_text('<!doctype html><html lang="ko"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>The Anatomy of Scent — Midterm Presentation</title><style>'+css+'</style></head><body><main id="stage">'+''.join(html)+'</main><nav class="controls" aria-label="슬라이드 조작"><button id="previous" aria-label="이전 슬라이드">←</button><span id="count"></span><button id="next" aria-label="다음 슬라이드">→</button><button id="print">PDF / Print</button></nav><script>'+js+'</script></body></html>')
    notes=['# 발표자 노트 — 약 5–6분','', '본 발표는 1–15장 중심으로 진행하고 16–21장은 부록으로 사용합니다. 14장에서 별도 HTML 프로토타입을 시연합니다.','']
    for i,s in enumerate(slides,1):notes += [f'## {i:02d}. {s["title"] or "The Anatomy of Scent"}',s['notes'],'']
    (ROOT/'SPEAKER-NOTES.md').write_text('\n'.join(notes))

def inline(t):
    t=escape(t);t=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',t);t=re.sub(r'`(.+?)`',r'<code>\1</code>',t);t=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'<a href="\2">\1</a>',t);return t

def proposal_exports():
    lines=(ROOT/'PROJECT-PROPOSAL.md').read_text().splitlines();parts=[];in_table=False;in_pre=False
    doc=Document();sec=doc.sections[0];sec.top_margin=DocInches(.75);sec.bottom_margin=DocInches(.75);sec.left_margin=DocInches(.85);sec.right_margin=DocInches(.85)
    for style in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[style];st.font.name=SANS;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),SANS);st.font.color.rgb=__import__('docx').shared.RGBColor.from_string(INK[1:])
    doc.styles['Normal'].font.size=DocPt(10);doc.styles['Normal'].paragraph_format.line_spacing=1.45;doc.styles['Normal'].paragraph_format.space_after=DocPt(7)
    footer=sec.footer.paragraphs[0];footer.text='THE ANATOMY OF SCENT / MIDTERM DESIGN PROPOSAL';footer.style='Caption'
    table_rows=[]
    def flush_table():
        nonlocal table_rows
        if table_rows:
            table=doc.add_table(rows=0,cols=len(table_rows[0]));table.style='Light Shading Accent 1'
            for row in table_rows:
                cells=table.add_row().cells
                for cell,t in zip(cells,row):cell.text=t
            table_rows=[]
    for line in lines:
        if line.startswith('```'):
            if in_pre:parts.append('</pre>')
            else:parts.append('<pre>')
            in_pre=not in_pre;continue
        if in_pre:parts.append(escape(line)+'\n');doc.add_paragraph(line,style='No Spacing');continue
        if line.startswith('|'):
            cols=[x.strip() for x in line.strip('|').split('|')]
            if all(re.fullmatch(r'[- :]+',x) for x in cols):continue
            if not in_table:parts.append('<table>');in_table=True
            parts.append('<tr>'+''.join('<td>'+inline(x)+'</td>' for x in cols)+'</tr>');table_rows.append(cols);continue
        if in_table:parts.append('</table>');in_table=False;flush_table()
        if not line.strip():continue
        if line.startswith('!['):
            m=re.search(r'!\[([^\]]*)\]\(([^)]+)\)',line)
            if m:parts.append(f'<figure><img src="{escape(m[2])}" alt="{escape(m[1])}"><figcaption>자체 생성 시각 자료 / 감각의 비유</figcaption></figure>');doc.add_picture(str(ROOT/m[2]),width=DocInches(6.0))
            continue
        match=re.match(r'^(#{1,3}) (.+)',line)
        if match:
            level=len(match[1]);parts.append(f'<h{level}>{inline(match[2])}</h{level}>');doc.add_heading(match[2],level=0 if level==1 else level-1);continue
        clean=re.sub(r'\*\*|`','',line);clean=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',clean)
        if line.startswith('- '):parts.append('<p class="item">'+inline(line[2:])+'</p>');doc.add_paragraph(clean[2:],style='List Bullet')
        elif re.match(r'^\d+\. ',line):parts.append('<p class="item">'+inline(line)+'</p>');doc.add_paragraph(clean)
        else:parts.append('<p>'+inline(line)+'</p>');doc.add_paragraph(clean)
    if in_table:parts.append('</table>');flush_table()
    doc.core_properties.title='The Anatomy of Scent — 프로젝트 기획안';doc.core_properties.subject='Week 06 Midterm Design Presentation';doc.save(ROOT/'Project-Proposal.docx')
    css='''@page{size:A4;margin:18mm 19mm 20mm}body{font-family:"Noto Sans CJK KR",Arial,sans-serif;color:#263d32;font-size:10.5pt;line-height:1.8;margin:0}h1{font-family:Georgia,serif;font-size:36pt;line-height:1.1;letter-spacing:-1px;margin:0 0 8pt}h2{font-size:19pt;line-height:1.4;margin:28pt 0 12pt;break-after:avoid}h3{font-size:12pt;color:#657369;margin:8pt 0 23pt}p{margin:0 0 10pt;word-break:keep-all;orphans:3;widows:3}table{width:100%;border-collapse:collapse;font-size:8.8pt;margin:12pt 0 16pt;break-inside:auto}tr{break-inside:avoid}td{border-bottom:1px solid #d7dacf;padding:8pt 7pt;vertical-align:top}tr:first-child{background:#eaf0dc;font-weight:600}figure{margin:18pt 0;break-inside:avoid}figure img{width:100%;height:auto}figcaption{font-size:8pt;color:#657369;margin-top:6pt}.item{padding-left:12pt;border-left:2px solid #e0e6d5}pre{font-family:"Noto Sans Mono CJK KR",monospace;font-size:9pt;line-height:1.6;white-space:pre-wrap;background:#f1f2eb;padding:15pt;break-inside:avoid}a{color:#385447;text-decoration:underline}code{font-family:inherit;background:#edf0e4;padding:1pt 3pt}strong{font-weight:600}'''
    (ROOT/'project-proposal.html').write_text('<!doctype html><html lang="ko"><head><meta charset="UTF-8"><title>The Anatomy of Scent — Project Proposal</title><style>'+css+'</style></head><body>'+''.join(parts)+'</body></html>')

if __name__=='__main__':
    diagrams();make_slides();proposal_exports()
    # Slide export runs after SVGs are rendered by export.mjs.
    if all((ASSETS/f'{n}.png').exists() for n in ['sitemap','user-flow','wireframe-desktop','wireframe-perspectives','wireframe-compare','wireframe-responsive']) and (ASSETS/'study-feel.png').exists():
        export_slides();print('Generated 21-slide PPTX, HTML viewer, Word proposal and source documents.')
    else:print('Diagram sources and Word proposal generated. Run node submission/source/export.mjs to complete rendering.')
